"""Build 00_Requisition/Science_Requisition_Weeks_1_to_10.xlsx from the extracted kit lists.

kit_lists.json: every item and class quantity by week, extracted from the kit cards (Section 5) of every lesson file and checked.
canonical.json: the same items matched across year groups, with draft Chinese names.
Run: python3 00_Requisition/source/build_req.py
If a kit card changes, update kit_lists.json (or re-extract) and run this again.
"""
import json, re, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

import os
S=os.path.join(os.path.dirname(os.path.abspath(__file__)),'')
kit=json.load(open(S+'kit_lists.json'))
can=json.load(open(S+'canonical.json'))
OVR={'device_phone_tablet':'collect_recycle','headphones':'collect_recycle','paper_towels_lab_dispenser':'prep_room_equipment',
     'booklet_logbook_y7':'print_job','booklet_dictionary_y7':'print_job'}
KEPT={'name_cards','group_role_cards_y9','folder_a3_group','envelope_c5'}
for c in can['canonical']:
    if c['key'] in OVR: c['category']=OVR[c['key']]
YEARS=['Y7','Y8','Y9']; YN={'Y7':'Year 7','Y8':'Year 8','Y9':'Year 9'}
F='Arial'
HF=Font(name=F,bold=True,color='FFFFFF',size=10); HFILL=PatternFill('solid',fgColor='00235F')
NF=Font(name=F,size=10); BF=Font(name=F,size=10,bold=True)
INFILL=PatternFill('solid',fgColor='FFFF00'); SECFILL=PatternFill('solid',fgColor='E0EFF8')
thin=Side(style='thin',color='B7C2D3'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
WRAP=Alignment(wrap_text=True,vertical='top'); CEN=Alignment(horizontal='center',vertical='top',wrap_text=True)

# ---------- unit normalisation ----------
GEN={'sheet':'sheets','envelope':'envelopes','roll':'rolls','clipboard':'clipboards','calculator':'calculators','kettle':'kettles','jug':'jugs','spoon':'teaspoons','spoons':'teaspoons','saucer':'saucers','lids':'saucers','cloth':'cloths','flask':'flasks','box':'boxes','cup':'cups','stopwatch':'stopwatches','timer':'stopwatches','timers':'stopwatches','cylinder':'cylinders','tray':'trays','bucket':'buckets','syringe':'syringes','book':'books','bottle':'bottles','pair':'pairs','jugs or cylinders':'jugs','metre sticks or tape measures':'metre sticks','packs or sticks':'glue sticks','hanger strings':'lengths','pieces 20 x 10 cm':'pieces','space overnight':'drawer','pen':'pens','block':'blocks'}
CONV={('sticky_notes','pads'):('notes',100),('kitchen_roll','sheets'):('rolls',1/60),('kitchen_roll','group supplies'):('rolls',0.1),
      ('gloves_disposable','box (shared)'):('pairs',10),('ice_cubes','cubes'):('kg',0.025),('flour_plain','teaspoons'):('g',3),
      ('salt_table','kg'):('g',1000),('salt_table','teaspoons'):('g',6)}
def norm_unit(key,unit,qty):
    u=unit.strip()
    for pre in ('A4 ','A3 '):
        if u.startswith(pre): u=u[len(pre):]
    if (key,u) in CONV:
        nu,f=CONV[(key,u)]; return nu, round(qty*f,3)
    if key=='marker_permanent_black' and u in ('pens','pen'): return 'markers',qty
    return GEN.get(u,u), qty

# ---------- build grid rows: (key, plan, unit) ----------
rows={}   # (key,plan,unit) -> dict
meta={}
for c in can['canonical']:
    meta[c['key']]=c
    for m in c['members']:
        it=kit[m['year']]['items'][m['index']]
        for u in it['uses']:
            unit,qty=norm_unit(c['key'],u['unit'],u['qty'])
            u=dict(u,qty=qty)
            k=(c['key'],it['plan'],unit)
            r=rows.setdefault(k,{'q':collections.defaultdict(float),'notes':[],'safety':set(),'src':set()})
            cell=(m['year'],u['week'])
            if c['reusable']: r['q'][cell]=max(r['q'][cell],u['qty'])
            else: r['q'][cell]+=u['qty']
            if u['note']: r['notes'].append(f"{m['year']} W{u['week']}: {u['note']}")
        if it['safety_note']:
            for k2 in [k for k in rows if k[0]==c['key']]: rows[k2]['safety'].add(it['safety_note'])
        for k2 in [k for k in rows if k[0]==c['key'] and k[1]==it['plan']]: rows[k2]['src'].add(it['source'])
CATORDER=['buy_consumable','buy_equipment','stationery_card_print','prep_room_equipment','canteen_or_kitchen','collect_recycle','teacher_made','print_job']
CATNAME={'buy_consumable':'Buy: used up','buy_equipment':'Buy: reusable','stationery_card_print':'Stationery and card','prep_room_equipment':'Prep room equipment','canteen_or_kitchen':'Canteen and kitchen','collect_recycle':'Collect or borrow','teacher_made':'Teacher made','print_job':'Print job (reprographics)'}
PLANNAME={'both':'Both plans','A':'Plan A (freezer)','B':'Plan B (no freezer)'}
def kind(key):
    m=meta[key]
    return 'kept' if key in KEPT else ('reus' if m['reusable'] else 'used')
TYPENAME={'kept':'Kept by class','reus':'Reusable','used':'Used up'}
def firstweek(r): return min(w for (y,w) in r['q'])
keys=sorted(rows, key=lambda k:(CATORDER.index(meta[k[0]]['category']), firstweek(rows[k]), meta[k[0]]['item_en'].lower(), k[1]))

wb=Workbook()
# ---------- Read me ----------
rm=wb.active; rm.title='Read me'
# ---------- Grid ----------
gd=wb.create_sheet('Data by week')
hdr=['Line','Item','中文名称（草稿）','Category','Plan','Type','Unit','Counts?']
for y in YEARS:
    for w in range(1,11): hdr.append(f'{y} W{w}')
for w in range(1,11): hdr.append(f'All W{w}')
hdr+= ['Y7 total','Y8 total','Y9 total']
for i,h in enumerate(hdr,1):
    c=gd.cell(1,i,h); c.font=HF; c.fill=HFILL; c.alignment=CEN; c.border=BOX
COL={h:i for i,h in enumerate(hdr,1)}
GR={}
for n,k in enumerate(keys, start=2):
    key,plan,unit=k; m=meta[key]; r=rows[k]
    line=f'{key}|{plan}|{unit}'
    vals=[line,m['item_en'],m['item_zh'],CATNAME[m['category']],PLANNAME[plan],TYPENAME[kind(key)],unit]
    for i,v in enumerate(vals,1): gd.cell(n,i,v).font=NF
    pc=f'$E{n}'
    gd.cell(n,COL['Counts?'],f'=IF(OR({pc}="Both plans",\'Read me\'!$C$9="A and B",AND({pc}="Plan A (freezer)",\'Read me\'!$C$9="A"),AND({pc}="Plan B (no freezer)",\'Read me\'!$C$9="B")),1,0)').font=NF
    for y in YEARS:
        for w in range(1,11):
            v=r['q'].get((y,w))
            cc=gd.cell(n,COL[f'{y} W{w}'],v if v else None); cc.font=NF
    for w in range(1,11):
        refs='+'.join(f'N({L(COL[f"{y} W{w}"])}{n})' for y in YEARS)
        gd.cell(n,COL[f'All W{w}'],f'={refs}').font=NF
    for y in YEARS:
        a=L(COL[f'{y} W1']); b=L(COL[f'{y} W10'])
        gd.cell(n,COL[f'{y} total'],f'=SUM({a}{n}:{b}{n})').font=NF
    GR[k]=n
gd.freeze_panes='C2'; gd.auto_filter.ref=f'A1:{L(len(hdr))}{len(keys)+1}'
gd.column_dimensions['A'].width=28; gd.column_dimensions['B'].width=34; gd.column_dimensions['C'].width=20
for h in ['Category','Plan','Type','Unit']: gd.column_dimensions[L(COL[h])].width=16
for i in range(COL['Counts?'],len(hdr)+1): gd.column_dimensions[L(i)].width=7.5
NROW=len(keys)+1
G="'Data by week'"
def rng(h): return f"{G}!${L(COL[h])}$2:${L(COL[h])}${NROW}"
LINES=rng('Line'); CNT=rng('Counts?')

# group grid rows by (key, unit) for display sheets (plans combined via Counts?)
groups=collections.OrderedDict()
for k in keys: groups.setdefault((k[0],k[2]),[]).append(k)

def weeks_used(gk):
    ws=collections.defaultdict(set)
    for k in groups[gk]:
        for (y,w) in rows[k]['q']: ws[y].add(w)
    return '; '.join(f"{YN[y]}: W{', W'.join(str(x) for x in sorted(ws[y]))}" for y in YEARS if ws[y])
def plans(gk):
    ps=sorted(set(k[1] for k in groups[gk]),key=lambda p:['both','A','B'].index(p))
    return ', '.join(PLANNAME[p] for p in ps)
def notes(gk):
    s=set(); n=[]
    for k in groups[gk]: s|=rows[k]['safety']; n+=rows[k]['notes']
    return ' '.join(sorted(s)), ' | '.join(n)
def sumifs(colh, gk):
    # sum over grid rows of this group with Counts?=1
    parts=[f'SUMPRODUCT(({LINES}="{k[0]}|{k[1]}|{k[2]}")*{CNT}*N(+{rng(colh)}))' for k in groups[gk]]
    return '+'.join(parts)
def rowref(colh,k): return f"{G}!${L(COL[colh])}${GR[k]}"
def cntref(k): return f"{G}!${L(COL['Counts?'])}${GR[k]}"
def year_total(gk,y,reusable):
    if not reusable:
        return '+'.join(f'{rowref(f"{y} total",k)}*{cntref(k)}' for k in groups[gk])
    return 'MAX(' + ','.join('+'.join(f'N({rowref(f"{y} W{w}",k)})*{cntref(k)}' for k in groups[gk]) for w in range(1,11)) + ')'
def peak_all(gk):
    return 'MAX(' + ','.join('+'.join(f'{rowref(f"All W{w}",k)}*{cntref(k)}' for k in groups[gk]) for w in range(1,11)) + ')'

def header(ws, cols, widths):
    for i,(h,wd) in enumerate(zip(cols,widths),1):
        c=ws.cell(1,i,h); c.font=HF; c.fill=HFILL; c.alignment=CEN; c.border=BOX
        ws.column_dimensions[L(i)].width=wd
    ws.freeze_panes='C2'; ws.row_dimensions[1].height=32

def write_row(ws,n,vals):
    for i,v in enumerate(vals,1):
        c=ws.cell(n,i,v); c.font=NF; c.alignment=WRAP; c.border=BOX

CL="'Read me'!$C$5"; C8="'Read me'!$C$6"; C9="'Read me'!$C$7"
CLS={'Y7':CL,'Y8':C8,'Y9':C9}

# ---------- Order list ----------
ol=wb.create_sheet('Order list',1)
cols=['Item','中文名称（草稿）','Specification','Category','Type','Year 7 per class','Year 8 per class','Year 9 per class','Quantity to order','Unit','First needed','Weeks used','Plan','Safety and notes','Ordered','Received']
header(ol,cols,[30,18,34,16,10,10,10,10,11,12,9,30,16,40,9,9])
n=2
for gk in groups:
    key,unit=gk; m=meta[key]
    if m['category'] not in ('buy_consumable','buy_equipment','stationery_card_print'): continue
    kd=kind(key); reus=kd!='used'; saf,nt=notes(gk)
    fw=min(firstweek(rows[k]) for k in groups[gk])
    yv=[f'=ROUND({year_total(gk,y,reus)},2)' for y in YEARS]
    if kd=='reus': tot=f'=MAX(F{n}:H{n})'
    else: tot='=ROUNDUP('+'+'.join(f'{L(6+i)}{n}*{CLS[y]}' for i,y in enumerate(YEARS))+',0)'
    write_row(ol,n,[m['item_en'],m['item_zh'],m['spec'],CATNAME[m['category']],TYPENAME[kd],*yv,tot,unit,f'Week {fw}',weeks_used(gk),plans(gk),saf,None,None])
    for i in (6,7,8,9): ol.cell(n,i).alignment=CEN
    ol.cell(n,9).font=BF
    for i in (15,16): ol.cell(n,i).fill=INFILL
    n+=1
ol.auto_filter.ref=f'A1:P{n-1}'; OL_N=n-1

# ---------- Prep room check ----------
pr=wb.create_sheet('Prep room check',2)
cols=['Item','中文名称（草稿）','Specification','Type','Most needed at once','Unit','First needed','Weeks used','Plan','Already have','Shortfall to order','Safety and notes']
header(pr,cols,[30,18,30,10,11,12,9,32,16,10,11,40])
n=2
for gk in groups:
    key,unit=gk; m=meta[key]
    if m['category']!='prep_room_equipment': continue
    kd=kind(key); saf,nt=notes(gk)
    fw=min(firstweek(rows[k]) for k in groups[gk])
    if kd=='reus': need='=MAX('+','.join(year_total(gk,y,True) for y in YEARS)+')'
    elif kd=='kept': need='=ROUNDUP('+'+'.join(f'({year_total(gk,y,True)})*{CLS[y]}' for y in YEARS)+',0)'
    else: need='=ROUNDUP('+'+'.join(f'({year_total(gk,y,False)})*{CLS[y]}' for y in YEARS)+',0)'
    write_row(pr,n,[m['item_en'],m['item_zh'],m['spec'],TYPENAME[kd],need,unit,f'Week {fw}',weeks_used(gk),plans(gk),None,f'=MAX(0,E{n}-N(J{n}))',saf])
    pr.cell(n,10).fill=INFILL; pr.cell(n,11).font=BF
    for i in (5,10,11): pr.cell(n,i).alignment=CEN
    n+=1
pr.auto_filter.ref=f'A1:L{n-1}'

# ---------- Canteen / collect / teacher made ----------
def simple_sheet(title,cat,pos):
    ws=wb.create_sheet(title,pos)
    cols=['Item','中文名称（草稿）','Specification','Type','Year 7 per class','Year 8 per class','Year 9 per class','Total','Unit','First needed','Weeks used','Plan','Safety and notes']
    header(ws,cols,[30,18,30,10,10,10,10,10,12,9,32,16,44])
    n=2
    for gk in groups:
        key,unit=gk; m=meta[key]
        if m['category']!=cat: continue
        kd=kind(key); reus=kd!='used'; saf,nt=notes(gk)
        fw=min(firstweek(rows[k]) for k in groups[gk])
        yv=[f'=ROUND({year_total(gk,y,reus)},2)' for y in YEARS]
        tot=f'=MAX(E{n}:G{n})' if kd=='reus' else '=ROUNDUP('+'+'.join(f'{L(5+i)}{n}*{CLS[y]}' for i,y in enumerate(YEARS))+',0)'
        write_row(ws,n,[m['item_en'],m['item_zh'],m['spec'],TYPENAME[kd],*yv,tot,unit,f'Week {fw}',weeks_used(gk),plans(gk),(saf+' '+('Notes: '+nt if cat=='teacher_made' else '')).strip()])
        for i in (5,6,7,8): ws.cell(n,i).alignment=CEN
        ws.cell(n,8).font=BF
        n+=1
    ws.auto_filter.ref=f'A1:M{max(n-1,1)}'
simple_sheet('Canteen and kitchen','canteen_or_kitchen',3)
simple_sheet('Collect or borrow','collect_recycle',4)
simple_sheet('Teacher made','teacher_made',5)

# ---------- By week ----------
bw=wb.create_sheet('By week',6)
cols=['Week','Year','Item','中文名称（草稿）','Quantity per class','Unit','Category','Plan','What it is for']
header(bw,cols,[7,8,30,18,10,12,18,16,70])
recs=[]
for c in can['canonical']:
    for mm in c['members']:
        it=kit[mm['year']]['items'][mm['index']]
        for u in it['uses']:
            recs.append((u['week'],mm['year'],CATORDER.index(c['category']),c['item_en'],c['item_zh'],u['qty'],u['unit'],CATNAME[c['category']],PLANNAME[it['plan']],u['note']))
recs.sort(key=lambda r:(r[0],r[1],r[2],r[3]))
n=2
for r in recs:
    write_row(bw,n,[f'Week {r[0]}',YN[r[1]],r[3],r[4],r[5],r[6],r[7],r[8],r[9]])
    bw.cell(n,5).alignment=CEN
    n+=1
bw.auto_filter.ref=f'A1:I{n-1}'


# ---------- Read me content ----------
T=Font(name=F,size=14,bold=True,color='00235F'); H2=Font(name=F,size=11,bold=True,color='00235F')
rm.column_dimensions['A'].width=3; rm.column_dimensions['B'].width=34; rm.column_dimensions['C'].width=16; rm.column_dimensions['D'].width=90
rm['B1']='Think Like a Scientist: science requisition, Weeks 1 to 10'; rm['B1'].font=T
rm['B2']='Years 7, 8 and 9. Quantities come from the kit card in every lesson file (Section 5), checked against the Teacher Guide prep calendar.'; rm['B2'].font=NF
rm['B4']='Inputs (yellow cells: change these)'; rm['B4'].font=H2
inputs=[(5,'Year 7 classes',1,'Number of Year 7 classes taking the programme. Used-up items are multiplied by this.'),
        (6,'Year 8 classes',1,'Number of Year 8 classes. Used-up items are multiplied by this.'),
        (7,'Year 9 classes',1,'Number of Year 9 classes. Type 0 if Year 9 is not running this term.')]
for r,lab,v,note in inputs:
    rm.cell(r,2,lab).font=BF; c=rm.cell(r,3,v); c.font=BF; c.fill=INFILL; c.alignment=CEN; c.border=BOX; rm.cell(r,4,note).font=NF
rm.cell(9,2,'Freezer plan').font=BF
c=rm.cell(9,3,'A and B'); c.font=BF; c.fill=INFILL; c.alignment=CEN; c.border=BOX
rm.cell(9,4,'Type A (freezer available), B (no freezer) or "A and B". Decide by Week 4. Leave "A and B" until then so both are costed; no Plan A or Plan B item is needed before Week 5.').font=NF
dv=DataValidation(type='list',formula1='"A and B,A,B"',allow_blank=False); rm.add_data_validation(dv); dv.add('C9')
for r in (5,6,7):
    dv2=DataValidation(type='whole',operator='between',formula1='0',formula2='10'); rm.add_data_validation(dv2); dv2.add(f'C{r}')
rm.cell(10,2,'Other cells to fill in').font=BF
rm.cell(10,4,'"Already have" on the Prep room check sheet (for example, type 6 if the prep room has 6 stopwatches; the shortfall updates), and the "Ordered" and "Received" columns on the Order list.').font=NF
r=12
def para(head, lines):
    global r
    rm.cell(r,2,head).font=H2; r+=1
    for t in lines:
        c=rm.cell(r,2,'• '+t); c.font=NF; c.alignment=Alignment(wrap_text=True,vertical='top'); rm.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
        rm.row_dimensions[r].height=30 if len(t)>120 else 16
        r+=1
    r+=1
para('How the quantities work',[
 'Quantities are for one class of up to 24 students in 6 groups, including the teacher demonstration and the spares the kit cards name.',
 'Used-up items (salt, filter paper, card, tape) are added up across all ten weeks and multiplied by the number of classes, then rounded up.',
 'Reusable items (stopwatches, measuring cylinders, goggles) show the most any one class needs at once. The year groups are taught at different times, so they share them. If two classes are ever taught at the same time, add their columns together.',
 'Kept by class items (name cards, group folders, Year 9 role cards, Year 7 spinner envelopes) belong to one class for the series, so they are multiplied by the number of classes.',
 'Weeks, not dates. "First needed" is the week an item must be in school; order it at least a week earlier.',
 'Print jobs are not on the order list: worksheets, the Scientist\'s Logbook and the My Science Dictionary booklets go to reprographics (they are listed on the By week sheet). Card, laminating pouches, A3 and colour prints are included because they are stock to buy.',
 'Weeks 7 and 8 depend on what each group chooses (Year 7 test plans, Year 8 test options, Year 9 claims). The list costs the maximum the kit cards allow; trim it once the choices are made in Week 7.',
 'Chinese names are drafts for the purchasing team. Ask a colleague to check them before sending.'])
para('The sheets',[
 'Order list: everything to buy (consumables, reusable items a science prep room would not usually have, stationery and card). This is the requisition.',
 'Prep room check: standard school science equipment. Ask the technician to fill in "Already have"; add any shortfall to the requisition.',
 'Canteen and kitchen: food, ice, freezer space and a kettle to request from the canteen or boarding house. Check the allergy notes.',
 'Collect or borrow: things students and staff can bring in or lend (bottles, jars, newspaper, towels, textbooks, the teacher\'s phone).',
 'Teacher made: things the teacher prepares from the items above (mystery cups, indicator, evidence bags). Their ingredients are already on the other sheets.',
 'By week: every item, week by week, for staging kit before each lesson.',
 'Data by week: the numbers behind the totals. Edit a quantity here and every total updates.'])
para('Decide before ordering',[
 'Freezer: Plan A or Plan B, by Week 4 (cell C9). Can the canteen supply the ice on the day (Year 7 Week 6, Year 9 Weeks 6 and 8)?',
 'Laminator and A3 printing: does the school have both? If not, print the laminated cards on card instead.',
 'Thermometers: count them in Week 1. Year 7 Plan B needs 7 in Week 6; Year 9 needs 12 plus spares in Week 6.',
 'Year 8 felt-tip pens: the list buys 6 full sets of pens A to D (24 pens), one set per group. Buying only 2 of each brand works if the teacher pre-dots every strip (adds about 10 minutes of prep). Run the pen pre-test in Week 1.',
 'Allergies: check for wheat, flour and corn allergies before Year 7 Plan B (Weeks 5, 7, 8) and Year 8 Weeks 4 and 8.',
 'Year 8 soil test (Week 8): check that the four soil locations named in the Case Bible exist on this campus and are safe to dig.',
 'Not on the kit cards but likely useful: reusable adhesive (Blu Tack) for posters, permanent markers and labels, a guillotine for cutting card sets, badge holders for the Year 8 expert witnesses.'])

wb._sheets.append(wb._sheets.pop(wb._sheets.index(gd)))
for ws in wb.worksheets:
    ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
    if ws.title!='Read me': ws.print_title_rows='1:1'
from openpyxl.workbook.properties import CalcProperties
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),'Science_Requisition_Weeks_1_to_10.xlsx'))
print('ok',len(keys),'grid rows;',OL_N-1,'order rows;',len(recs),'by-week rows')
