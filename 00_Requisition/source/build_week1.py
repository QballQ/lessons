"""Build Week_1_Resources_Y7_Y8.xlsx: everything needed for Year 7 and Year 8 Week 1, one list per year,
for checking item by item. Quantities are for one class of up to 24 students in 6 groups, checked against the
lesson files (Sections 1, 5 and 6), the skills check and survey files, the language support pack, the Case Bible
and the designed printables, and compared with Science_Requisition_Weeks_1_to_10.xlsx.
Run: python3 build_week1.py"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'Week_1_Resources_Y7_Y8.xlsx')

NAVY = '00235F'
FONT = 'Arial'
HEAD_FILL = PatternFill('solid', fgColor='E0EFF8')
GROUP_FILL = PatternFill('solid', fgColor=NAVY)
INPUT_FILL = PatternFill('solid', fgColor='FFF2A8')
FLAG_FILL = PatternFill('solid', fgColor='FDE6E8')
THIN = Side(style='thin', color='B7C2D3')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Columns: item, qty (one class), unit / how, per, paper, sheets per copy, for, optional, requisition, check
# per: 'student', 'group', 'class' multiply by the number of classes; 'room' and 'once' do not.
# requisition: 'Yes', 'Yes, see note', 'No: paper printing', 'No: add', 'No: tool', 'No: may already exist'
Y7 = [
    ('Printing (paper copies)', [
        ('Skills check paper, Year 7', 24, 'copies; 5 A4 pages, double-sided', 'student', 'A4 paper', 3,
         'Week 1 baseline, 15 minutes, done alone; marked, filed and kept until Week 10.', '', 'No: paper printing',
         'The skills check file says "two sides of A4", but the designed paper is 5 pages. Print the designed version: 3 sheets each.'),
        ('Interest survey', 24, 'half-A4 copies; 2 per A4 sheet, cut in half', 'student', 'A4 paper', 0.5,
         'Week 1 baseline survey, 3 minutes, straight after the skills check.', '', 'No: paper printing', ''),
        ('Worksheet: What do scientists do?', 24, 'copies; 4 A4 pages, double-sided', 'student', 'A4 paper', 2,
         'Observation or inference sort, footprint table, Stretch, Challenge and the exit ticket strip.', '', 'No: paper printing',
         'The exit ticket strip is cut off the foot of page 4. Printed double-sided, the cut also goes through the bottom of page 3 (My thinking box). Check before printing.'),
        ('Skills check mark scheme (teacher copy)', 1, 'copy; 2 A4 pages', 'class', 'A4 paper', 2,
         'Marking the skills checks out of 20 during Week 1.', '', 'No: paper printing', ''),
        ('Kit card (teacher copy)', 1, 'copy; A5', 'class', 'A5 paper', 1,
         'Kit list, prep, hazards and clean-up for the lesson.', '', 'No: paper printing',
         'It still says role cards "1 set of 5". Use the 6-card sets below.'),
        ('Group role cards with duties', 6, 'sets; 6 cards per A4 sheet, cut and laminated', 'group', 'A4 paper', 1,
         'One set per table: Lead Scientist, Measurer, Recorder, Safety Officer, Reporter, Language buddy. Kept all term.', '', 'No: paper printing',
         'Laminating pouches are on the requisition (6 A4).'),
        ('Classroom language cards', 30, 'cards; A5, two-sided, laminated', 'class', 'A5 paper', 1,
         '1 per student plus 6 spares; help phrases at every desk, kept all term.', '', 'No: paper printing',
         'Laminating pouches are on the requisition (30 A5).'),
        ('Language support posters (Instruction words key, Classroom language side B, Think in any language)', 3, 'posters; A3, 1 of each', 'room', 'A3 paper', 1,
         'Wall posters, up all term.', '', 'Yes',
         'If Year 7 and Year 8 share a room, one set is enough.'),
        ('Sentence frame poster, Ice Engineers', 1, 'poster; A4', 'room', 'A4 paper', 1,
         'Wall bank of the Year 7 sentence frames with Chinese.', 'Teacher Guide', 'No: add',
         'The Teacher Guide asks for it before Week 1; the Week 1 lesson file does not mention it.'),
        ('Sentence frame cards, Ice Engineers', 6, 'sets of 16 cards; 4 A6 cards per A4 sheet, cut', 'group', 'A4 paper', 4,
         'One set per table for rehearsing answers in English. Kept all term.', 'Teacher Guide', 'No: add',
         'As above. Paper or card is not stated.'),
        ('My Science Dictionary, Year 7 (A5 booklet)', 28, 'booklets; 24 pages, folded and stapled', 'class', 'A4 90 gsm (booklet)', 6,
         '1 per student plus 4 spares; handed out with the logbook, used every lesson.', '', 'Yes',
         'Print job for reprographics. 28 pages (7 sheets) if the entries need more room.'),
        ("Scientist's Logbook, Year 7 (A5 booklet)", 24, 'booklets; page count not given', 'student', 'Booklet (reprographics)', None,
         'First page started at the end of the lesson; kept in the class box.', '', 'Yes',
         'Page count and spares are not given anywhere.'),
    ]),
    ('Card, laminating and stationery', [
        ('Footprint puzzle cards (Stage 1, 2, 3) on card', 6, 'sets of 3 A4 landscape cards', 'group', 'A4 card', 3,
         'Reveal cards for the 10-minute footprint puzzle, one stage at a time. Reused.', '', 'Yes, see note',
         'Card 3 needs the teacher note printed on its back, so it is double-sided; the requisition says single-sided. Card or laminated paper is allowed.'),
        ('Laminating pouches, A4 (role cards)', 6, 'pouches', 'group', '', None,
         'Laminate the 6 role card sheets, then cut.', '', 'Yes', ''),
        ('Laminating pouches, A5 (classroom language cards)', 30, 'pouches', 'class', '', None,
         'Laminate the 30 classroom language cards.', '', 'Yes', ''),
        ('Laminating pouches, A4 (puzzle cards on paper instead of card)', 18, 'pouches', 'group', '', None,
         'Only if the puzzle cards are printed on paper and laminated instead of card.', 'Optional', 'No: add',
         'Mentioned only inside the card note on the requisition.'),
        ('Envelopes, C4 (to hold A4)', 18, 'envelopes', 'group', '', None,
         '3 per group, labelled Stage 1, 2, 3; Stage 2 and 3 envelopes stay on the teacher desk.', '', 'Yes', ''),
        ('Name cards', 24, 'cards', 'student', '', None,
         '1 per student, set out by the group seating plan.', '', 'Yes', 'No template or size is given.'),
        ('Pencils', 24, 'pencils, plus a few spares', 'student', '', None,
         'Skills check, survey, worksheet and logbook.', '', 'Yes', 'Spares for the teacher desk are not counted.'),
        ('Class box (plastic, with lid)', 1, 'box', 'class', '', None,
         'Logbooks are collected into it at the end of the lesson.', '', 'Yes', ''),
        ('Class assessment folder', 1, 'folder', 'class', '', None,
         'Marked skills checks are filed here until Week 10.', '', 'No: may already exist', ''),
    ]),
    ('Equipment and tools', [
        ('Laminator', 1, 'machine (access)', 'once', '', None,
         'Role cards and classroom language cards.', '', 'No: tool', ''),
        ('Guillotine or scissors', 1, 'guillotine or pair', 'once', '', None,
         'Cutting the surveys in half, the role cards and the frame cards.', '', 'No: tool', ''),
        ('Rulers', 6, 'rulers, 1 per group', 'group', '', None,
         'Card 1 asks students to "count and measure"; every card has a 10 cm scale bar.', 'Optional', 'No: add',
         'Not on the kit card. Students can estimate against the scale bar, so decide whether to put one on each table.'),
    ]),
]

Y8 = [
    ('Printing (paper copies)', [
        ('Skills check paper, Year 8', 24, 'copies; 5 A4 pages, double-sided', 'student', 'A4 paper', 3,
         'Week 1 baseline, 15 minutes, done alone; marked, filed and kept until Week 10.', '', 'No: paper printing',
         'The skills check file says "two sides of A4", but the designed paper is 5 pages. Print the designed version: 3 sheets each.'),
        ('Interest survey', 24, 'half-A4 copies; 2 per A4 sheet, cut in half', 'student', 'A4 paper', 0.5,
         'Week 1 baseline survey, 3 minutes, straight after the skills check.', '', 'No: paper printing', ''),
        ('Worksheet: Think like a detective', 24, 'copies; 3 A4 pages (pages 1 and 2 back to back, page 3 on its own)', 'student', 'A4 paper', 2,
         'Scene map task: Core, Stretch, Challenge, word bank and exit ticket strip. Goes into the case file.', '', 'No: paper printing',
         'The exit ticket strip is cut off and collected. Check it does not cut into anything on the back.'),
        ('Case briefing and sign-in sheet', 24, 'copies; 2 A4 pages, double-sided', 'student', 'A4 paper', 1,
         'What happened, the note, the five evidence items and the Lab 3 sign-in sheet. Second item in the case file.', '', 'No: paper printing',
         'The lesson file describes one page; the design has two.'),
        ('Scene map', 24, 'copies; A4 landscape', 'student', 'A4 paper', 1,
         'Plan of Lab 3 on a lettered grid for the Do task. Third item in the case file.', '', 'No: paper printing', ''),
        ('Scene map, A3 wall copy', 1, 'copy; A3', 'class', 'A3 paper', 1,
         'Wall copy; Band 1 students point to a square on it for the exit ticket.', '', 'Yes', ''),
        ('Suspect profile cards', 24, 'sheets; 4 A6 cards per A4 sheet', 'student', 'A4 paper', 1,
         'The four suspects, in the case file; used again in the Week 2 Do Now.', '', 'No: paper printing', ''),
        ('Skills check mark scheme (teacher copy)', 1, 'copy; 2 A4 pages', 'class', 'A4 paper', 2,
         'Marking the skills checks during Week 1.', '', 'No: paper printing', ''),
        ('Kit card (teacher copy)', 1, 'copy; A5', 'class', 'A5 paper', 1,
         'Kit list, prep, hazards and clean-up for the lesson.', '', 'No: paper printing', ''),
        ('Teacher Pre-test Record sheet', 1, 'copy (Week 3 file, Section 6)', 'once', 'A4 paper', 1,
         'Recording the pen A to D results during the pen pre-test.', '', 'No: paper printing', ''),
        ('Group role cards with duties', 6, 'sets; 6 cards per A4 sheet, cut and laminated', 'group', 'A4 paper', 1,
         'One set per table. Kept all term.', '', 'No: paper printing',
         'Laminating pouches are on the requisition (6 A4). The Year 8 kit card does not mention role cards.'),
        ('Classroom language cards', 30, 'cards; A5, two-sided, laminated', 'class', 'A5 paper', 1,
         '1 per student plus 6 spares; help phrases at every desk, kept all term.', '', 'No: paper printing',
         'Laminating pouches are on the requisition (30 A5).'),
        ('Language support posters (Instruction words key, Classroom language side B, Think in any language)', 3, 'posters; A3, 1 of each', 'room', 'A3 paper', 1,
         'Wall posters, up all term.', '', 'Yes',
         'If Year 7 and Year 8 share a room, one set is enough.'),
        ('Sentence frame poster, Science Detectives', 1, 'poster; A4', 'room', 'A4 paper', 1,
         'Wall bank of the Year 8 sentence frames with Chinese.', 'Teacher Guide', 'No: add',
         'The Teacher Guide asks for it before Week 1; the Week 1 lesson file does not mention it.'),
        ('Sentence frame cards, Science Detectives', 6, 'sets of 15 cards; 4 A6 cards per A4 sheet, cut', 'group', 'A4 paper', 4,
         'One set per table for rehearsing answers in English. Kept all term.', 'Teacher Guide', 'No: add',
         'As above. Paper or card is not stated.'),
        ('My Science Dictionary, Year 8 (A5 booklet)', 28, 'booklets; 28 pages, folded and stapled', 'class', 'A4 90 gsm (booklet)', 7,
         '1 per student plus 4 spares; handed out inside each case file.', '', 'No: add',
         'Missing from the Year 8 requisition (Year 7 has it). The written spec says 32 pages (8 sheets); the designed booklet is 28 pages.'),
        ("Scientist's Logbook, Year 8 (A5 booklet)", 24, 'booklets; page count not given', 'student', 'Booklet (reprographics)', None,
         'A line under "What I found" at the end of the lesson; kept in the class box.', '', 'No: add',
         'Missing from the Year 8 requisition (Year 7 has it). Page count and spares are not given.'),
    ]),
    ('Card, laminating and stationery', [
        ('Case file covers on card', 24, 'A4 portrait, single-sided', 'student', 'A4 card', 1,
         'Front of each student\'s case file: detective name, class, group and role.', '', 'Yes',
         'Part of the 49 sheets of white card on the requisition.'),
        ('Evidence logs on card', 24, 'A4 landscape, single-sided', 'student', 'A4 card', 1,
         'Evidence table and suspect tracker, kept in the case file all ten weeks.', '', 'Yes',
         'Part of the 49 sheets of white card.'),
        ('White card for the note prop', 1, 'A4 sheet, plain', 'once', 'A4 card', 1,
         'The handwritten note (see Make beforehand).', '', 'Yes', 'Part of the 49 sheets of white card.'),
        ('A4 folders for case files', 24, 'folders', 'student', '', None,
         'Holds the cover, briefing, scene map, suspect cards, evidence log, worksheet and dictionary all ten weeks.', '', 'Yes', ''),
        ('Laminating pouches, A4 (role cards)', 6, 'pouches', 'group', '', None,
         'Laminate the 6 role card sheets, then cut.', '', 'Yes', ''),
        ('Laminating pouches, A5 (classroom language cards)', 30, 'pouches', 'class', '', None,
         'Laminate the 30 classroom language cards.', '', 'Yes', ''),
        ('Name cards', 24, 'cards', 'student', '', None,
         'Groups of 4 set out with name cards.', '', 'Yes', 'The kit card gives no number; 1 per student assumed.'),
        ('Pencils', 24, 'pencils', 'student', '', None,
         'Skills check, survey, worksheet and exit ticket.', '', 'Yes', ''),
        ('Masking tape', 1, 'roll', 'once', '', None,
         'Cover the brand name on each pen and write the letters A to D on the tape.', '', 'Yes', ''),
        ('Class box (plastic, with lid)', 1, 'box', 'class', '', None,
         'Case files and logbooks stay in it between lessons.', '', 'Yes', ''),
        ('Class assessment folder', 1, 'folder', 'class', '', None,
         'Marked skills checks are filed here until Week 10.', '', 'No: may already exist', ''),
    ]),
    ('Equipment and materials', [
        ('Black washable felt-tip pens, 4 different brands, 2 of each', 8, 'pens', 'once', '', None,
         'Pen pre-test (pens A to D, one spare of each); pen B writes the note. Used again in Week 3.', '', 'Yes',
         'Buy before Week 1.'),
        ('White coffee filters', 4, 'strips, 2 cm by 10 cm', 'once', '', None,
         'Chromatography strips for the pen pre-test, one per pen.', '', 'Yes',
         'Buy a pack: more strips are needed if a pen has to be swapped and tested again.'),
        ('Small clear cups', 4, 'cups', 'once', '', None,
         'Pen pre-test: water 1 cm deep, one cup per strip.', '', 'Yes', ''),
        ('Ruler (30 cm) and a pencil, for the teacher', 1, 'of each', 'once', '', None,
         'Marking the strips and checking the 24 cm footprint print.', '', 'No: add', ''),
        ('Laminator', 1, 'machine (access)', 'once', '', None,
         'Role cards and classroom language cards.', '', 'No: tool', ''),
        ('Guillotine or scissors', 1, 'guillotine or pair', 'once', '', None,
         'Cutting the filter strips, surveys, role cards, frame cards and exit ticket strips.', '', 'No: tool', ''),
    ]),
    ('Optional scene table (5 minutes to set up)', [
        ('Drinking glass, clear', 1, 'glass', 'once', '', None, 'Copy of evidence 2.', 'Optional', 'Yes', ''),
        ('Baking soda', 1, 'teaspoon, plus a pinch', 'class', '', None, 'Copy of the white powder (evidence 3), and part of the scene liquid.', 'Optional', 'Yes', ''),
        ('Dark tray', 1, 'tray', 'once', '', None, 'Holds the teaspoon of baking soda.', 'Optional', 'Yes, see note', 'The requisition lists a flat kit tray; the Case Bible asks for a dark one.'),
        ('Saucer', 1, 'saucer', 'once', '', None, 'Holds the scene liquid (evidence 4).', 'Optional', 'Yes', ''),
        ('Plain bar soap', 1, 'teaspoon, grated', 'class', '', None, 'Soap solution for the scene liquid (not washing-up liquid).', 'Optional', 'Yes', ''),
        ('Grater, teaspoon, heatproof jug (200 ml) and stirrer', 1, 'set', 'once', '', None, 'Making the soap and baking soda solutions.', 'Optional', 'No: tool', ''),
        ('Footprint print, true size (24 cm)', 1, 'print', 'once', 'A4 paper', 1, 'Paper footprint for the scene table; check it prints at 24 cm.', 'Optional', 'Yes', ''),
        ('Scene table sign: EVIDENCE: DO NOT TOUCH', 1, 'sign; A5', 'once', 'A5 paper', 1, 'Look-only sign.', 'Optional', 'No: paper printing', ''),
    ]),
    ('Make beforehand', [
        ('Pen pre-test strips (answer key)', 4, 'strips, one per pen', 'once', '', None,
         'About 20 minutes. Picks pen B for the note and gives the Week 3 answer key.', '', 'Yes',
         'Due two weeks before Week 3, so in the Week 1 week. Record the results on the Pre-test Record sheet.'),
        ('The note prop', 1, 'note', 'once', '', None,
         'Written with pen B on white card after the pre-test; on the scene table and held up again in Week 3.', '', 'Yes',
         'Slide 5 already shows a drawn note, so a photo is only needed if you want to replace it.'),
        ('Scene liquid (soap solution and baking soda solution mixed)', 1, 'saucer', 'class', '', None,
         'Copy of the puddle (evidence 4) for the scene table.', 'Optional', 'Yes', ''),
    ]),
]

COLS = ['Item', 'Quantity (one class)', 'Unit or how to print', 'Per', 'Total for your classes', 'Paper or card',
        'Sheets per copy', 'Sheets in total', 'What it is for', 'Optional?', 'On the requisition?', 'Check', 'Got it']
WIDTHS = [34, 10, 26, 10, 11, 15, 8, 9, 38, 12, 17, 44, 7]


def style(c, bold=False, color=None, fill=None, wrap=True, size=10, align='left'):
    c.font = Font(name=FONT, size=size, bold=bold, color=color or NAVY)
    c.alignment = Alignment(wrap_text=wrap, vertical='top', horizontal=align)
    if fill: c.fill = fill


def sheet(wb, title, year_label, groups):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws['A1'] = f'{year_label}, Week 1: resources needed'
    style(ws['A1'], bold=True, size=14, wrap=False)
    ws['A2'] = 'Quantities are for one class of up to 24 students in 6 groups. Type the number of classes in the yellow cell; totals update.'
    style(ws['A2'], wrap=False)
    ws['A3'] = 'Number of classes'
    style(ws['A3'], bold=True, wrap=False)
    ws['B3'] = 1
    style(ws['B3'], bold=True, fill=INPUT_FILL, align='center')
    ws['B3'].border = BOX
    ws['C3'] = 'Items "per student", "per group" or "per class" are multiplied by this; "per room" and "once" are not.'
    style(ws['C3'], wrap=False, size=9)
    r = 5
    for i, h in enumerate(COLS, 1):
        c = ws.cell(row=r, column=i, value=h)
        style(c, bold=True, fill=HEAD_FILL)
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[i - 1]
    ws.freeze_panes = 'B6'
    first_data = r + 1
    for gname, rows in groups:
        r += 1
        c = ws.cell(row=r, column=1, value=gname)
        for col in range(1, len(COLS) + 1):
            style(ws.cell(row=r, column=col), bold=True, color='FFFFFF', fill=GROUP_FILL, wrap=False)
        for (item, qty, unit, per, paper, spc, purpose, opt, req, check) in rows:
            r += 1
            vals = [item, qty, unit, 'per ' + per, None, paper or None, spc, None, purpose, opt or None, req, check or None, None]
            for col, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=col, value=v)
                style(c, align='center' if col in (2, 4, 5, 7, 8, 13) else 'left')
                c.border = BOX
            ws.cell(row=r, column=5, value=f'=IF(OR(D{r}="per student",D{r}="per group",D{r}="per class"),B{r}*$B$3,B{r})')
            if spc is not None:
                ws.cell(row=r, column=8, value=f'=ROUNDUP(E{r}*G{r},0)')
            if req.startswith('No: add') or req == 'Yes, see note':
                ws.cell(row=r, column=11).fill = FLAG_FILL
            if check:
                ws.cell(row=r, column=12).fill = FLAG_FILL
    last = r
    # paper totals
    r += 2
    ws.cell(row=r, column=1, value='Paper and card for printing (sheets in total)')
    style(ws.cell(row=r, column=1), bold=True, wrap=False)
    for kind in ['A4 paper', 'A3 paper', 'A5 paper', 'A4 card', 'A4 90 gsm (booklet)']:
        r += 1
        ws.cell(row=r, column=1, value=kind)
        style(ws.cell(row=r, column=1), wrap=False)
        ws.cell(row=r, column=2, value=f'=SUMIFS($H${first_data}:$H${last},$F${first_data}:$F${last},A{r})')
        style(ws.cell(row=r, column=2), bold=True, align='center')
        ws.cell(row=r, column=2).border = BOX
        ws.cell(row=r, column=3, value='sheets')
        style(ws.cell(row=r, column=3), wrap=False)
    r += 1
    ws.cell(row=r, column=1, value='Booklets of unknown length (logbooks) are not counted.')
    style(ws.cell(row=r, column=1), size=9, wrap=False)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3  # wide list: A3 landscape keeps the text readable
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '5:5'
    return ws


def readme(wb):
    ws = wb.active
    ws.title = 'Read me'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 120
    lines = [
        ('Week 1 resources: Year 7 and Year 8', True, 14),
        ('Everything needed to teach Week 1, one list per year, so it can be checked item by item. Year 9 is left out.', False, 10),
        ('It is not split by where things come from (order, prep room, canteen). Each list is grouped by kind: printing, card and stationery, equipment, and (Year 8) the optional scene table and things to make beforehand.', False, 10),
        ('Quantities are for one class of up to 24 students in 6 groups. Type the number of classes in the yellow cell at the top of each list. The "Got it" column is for your ticks.', False, 10),
        ('Checked against the lesson files (lesson plan, kit card and printables), the skills check and survey files, the language support pack, the Year 8 Case Bible and the designed printables, and compared with the full requisition workbook. Pink cells are the ones to look at.', False, 10),
        ('', False, 10),
        ('Things to decide or fix', True, 12),
        ('1. Year 8 requisition is missing the My Science Dictionary (28 booklets) and the Scientist\'s Logbook (24). Year 7 has both. Added to the Year 8 list.', False, 10),
        ('2. Skills check papers: the files say "two sides of A4", but the designed papers are 5 pages. Print the designed version: 3 sheets each, 72 sheets per class.', False, 10),
        ('3. Ordinary paper printing (worksheets, skills checks, surveys, case file pages, role cards and so on) is not on the requisition. It is listed here with sheet counts for reprographics.', False, 10),
        ('4. Year 7 footprint puzzle: Card 3 has the teacher note on its back, so print it double-sided. The requisition says single-sided.', False, 10),
        ('5. Year 7 kit card says role cards "1 set of 5". The language support pack uses 6-card sets (with Language buddy), which is what is listed.', False, 10),
        ('6. Sentence frame posters and cards: the Teacher Guide says to print them before Week 1, but they are not on the requisition or in the Week 1 lesson files.', False, 10),
        ('7. Year 8 pen pre-test: buy the pens before Week 1 and do the 20-minute pre-test in the Week 1 week (two weeks before Week 3). The note prop is written with pen B after the pre-test.', False, 10),
        ('8. If Year 7 and Year 8 are taught in the same room, one set of the three A3 language posters is enough.', False, 10),
        ('9. Not listed because they are normal classroom kit: projector and screen, the slides, a clock or timer.', False, 10),
    ]
    for i, (t, b, s) in enumerate(lines, 1):
        c = ws.cell(row=i + 1, column=2, value=t)
        style(c, bold=b, size=s)
    ws['B4'].alignment = Alignment(wrap_text=True, vertical='top')


wb = Workbook()
readme(wb)
sheet(wb, 'Year 7 Week 1', 'Year 7 Ice Engineers', Y7)
sheet(wb, 'Year 8 Week 1', 'Year 8 Science Detectives', Y8)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', os.path.abspath(OUT))
