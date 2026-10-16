"""Build Resources_Weeks_1_to_10_Y7_Y8.xlsx from lessons_week_data.json (one record per lesson, checked against the
lesson files, kit cards, printables, Teacher Guide and Case Bible, and compared with the requisition).

Sheets:
  Read me        what it is, how to use it, things to decide
  All weeks      every resource for every Year 7 and Year 8 lesson on one sheet, with the inputs (classes, freezer plan)
  Prep by week   every prep task with when, minutes and who; a week-by-week summary of the prep load
  Cut the prep   where the time goes and the options for reducing it
Run: python3 build_all_weeks.py"""
import json, os
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, 'lessons_week_data.json')
OUT = os.path.join(HERE, '..', 'Resources_Weeks_1_to_10_Y7_Y8.xlsx')

NAVY, FONT = '00235F', 'Arial'
HEAD = PatternFill('solid', fgColor='E0EFF8')
BAND = PatternFill('solid', fgColor=NAVY)
INPUT = PatternFill('solid', fgColor='FFF2A8')
FLAG = PatternFill('solid', fgColor='FDE6E8')
ODD = PatternFill('solid', fgColor='F6F8FB')
THIN = Side(style='thin', color='B7C2D3')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GROUP_ORDER = ['Printing (paper copies)', 'Card, laminating and stationery', 'Equipment and materials', 'Food and consumables', 'Make beforehand']
WHEN_ORDER = ['2 or more weeks before', 'The week before', 'The day before', 'On the day, before the lesson', 'Set-up by students in the lesson', 'After the lesson']
PAPERS = ['A4 paper', 'A3 paper', 'A5 paper', 'A4 card', 'A4 90 gsm (booklet)']


def st(c, bold=False, color=NAVY, fill=None, wrap=True, size=10, align='left'):
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    c.alignment = Alignment(wrap_text=wrap, vertical='top', horizontal=align)
    if fill: c.fill = fill


def header(ws, row, cols, widths):
    for i, (h, w) in enumerate(zip(cols, widths), 1):
        c = ws.cell(row=row, column=i, value=h)
        st(c, bold=True, fill=HEAD)
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = w


def year_of(L):
    return 'Year 7' if L['id'].startswith('Y7') else 'Year 8'


def build():
    lessons = sorted(json.load(open(DATA)), key=lambda L: (L['week'], L['id']))
    wb = Workbook()
    rd = wb.active
    rd.title = 'Read me'
    aw = wb.create_sheet('All weeks')
    pw = wb.create_sheet('Prep by week')
    cw = wb.create_sheet('Cut the prep')

    # ---------- All weeks ----------
    aw.sheet_view.showGridLines = False
    aw['A1'] = 'Years 7 and 8, Weeks 1 to 10: resources needed'
    st(aw['A1'], bold=True, size=14, wrap=False)
    aw['A2'] = 'Every lesson on one sheet. Quantities are for one class of up to 24 students in 6 groups. Change the yellow cells; totals update. Use the filter arrows to show one week, one year or one group.'
    st(aw['A2'], wrap=False)
    for r, (lab, val, note) in enumerate([
            ('Year 7 classes', 1, 'Items per student, per group or per class are multiplied by the classes in that year.'),
            ('Year 8 classes', 1, ''),
            ('Year 7 freezer plan', 'A and B', 'A (freezer), B (no freezer) or "A and B". Items for the plan not chosen count as 0. Decide by Week 4.')], start=3):
        aw.cell(row=r, column=1, value=lab); st(aw.cell(row=r, column=1), bold=True, wrap=False)
        c = aw.cell(row=r, column=3, value=val); st(c, bold=True, fill=INPUT, align='center'); c.border = BOX
        aw.cell(row=r, column=4, value=note); st(aw.cell(row=r, column=4), size=9, wrap=False)
    dv = DataValidation(type='list', formula1='"A,B,A and B"', allow_blank=False)
    aw.add_data_validation(dv); dv.add('C5')
    cols = ['Week', 'Year', 'Group', 'Item', 'Quantity (one class)', 'Unit or how to print', 'Per', 'Plan', 'Total for your classes',
            'Paper or card', 'Sheets per copy', 'Sheets in total', 'What it is for', 'Optional?', 'Already have (from an earlier week)',
            'On the requisition?', 'Check', 'Got it']
    widths = [6, 8, 18, 34, 9, 26, 10, 8, 10, 14, 7, 8, 36, 9, 20, 16, 40, 6]
    H = 8
    header(aw, H, cols, widths)
    aw.freeze_panes = aw.cell(row=H + 1, column=5)
    r = H
    for L in lessons:
        rows = sorted(L['resources'], key=lambda x: (GROUP_ORDER.index(x['group']) if x['group'] in GROUP_ORDER else 9))
        for x in rows:
            r += 1
            q = x.get('quantity')
            reuse = x.get('reuse') or ''
            vals = [L['week'], year_of(L), x['group'], x['item'], q if q is not None else None,
                    x.get('unit') or '' if q is not None else ((x.get('quantity_text') or 'not stated') + (('; ' + x['unit']) if x.get('unit') else '')),
                    'per ' + x['basis'], x['plan'] if x['plan'] != 'Both' else '', None, x.get('paper') or None,
                    x.get('sheets_per_copy'), None, x['purpose'], 'Optional' if x['optional'] else None, reuse or None,
                    x['in_requisition'], x.get('check') or None, None]
            for i, v in enumerate(vals, 1):
                c = aw.cell(row=r, column=i, value=v)
                st(c, align='center' if i in (1, 5, 7, 8, 9, 11, 12, 14, 18) else 'left')
                c.border = BOX
                if L['week'] % 2: c.fill = ODD
            aw.cell(row=r, column=9, value=(
                f'=IF(E{r}="","",IF(OR(AND(H{r}="Plan A",$C$5="B"),AND(H{r}="Plan B",$C$5="A")),0,'
                f'IF(OR(G{r}="per student",G{r}="per group",G{r}="per class"),E{r}*IF(B{r}="Year 7",$C$3,$C$4),E{r})))'))
            aw.cell(row=r, column=12, value=f'=IF(OR(K{r}="",O{r}<>"",I{r}=""),"",ROUNDUP(I{r}*K{r},0))')
            if x['in_requisition'] in ('No: add', 'Yes, see note'): aw.cell(row=r, column=16).fill = FLAG
            if x.get('check'): aw.cell(row=r, column=17).fill = FLAG
    last = r
    aw.auto_filter.ref = f'A{H}:R{last}'
    # paper totals box
    aw['J3'] = 'Printing, all weeks (sheets)'
    st(aw['J3'], bold=True, wrap=False)
    for i, p in enumerate(PAPERS):
        rr = 4 + i if i < 3 else 4 + i - 3
        cc = 10 if i < 3 else 13
        aw.cell(row=rr, column=cc, value=p); st(aw.cell(row=rr, column=cc), wrap=False, size=9)
        c = aw.cell(row=rr, column=cc + 2, value=f'=SUMIFS($L${H + 1}:$L${last},$J${H + 1}:$J${last},{get_column_letter(cc)}{rr})')
        st(c, bold=True, align='center'); c.border = BOX
    aw.page_setup.orientation = 'landscape'
    aw.page_setup.paperSize = aw.PAPERSIZE_A3
    aw.page_setup.fitToWidth = 1
    aw.page_setup.fitToHeight = 0
    aw.sheet_properties.pageSetUpPr.fitToPage = True
    aw.print_title_rows = f'{H}:{H}'

    # ---------- Prep by week ----------
    pw.sheet_view.showGridLines = False
    pw['A1'] = 'Prep by week: how much time each lesson takes to get ready'
    st(pw['A1'], bold=True, size=14, wrap=False)
    pw['A2'] = ('Minutes come from the kit cards where they give them; the rest are estimates (marked). "After" is marking, the tracker and putting kit away. '
                '"Could be done by someone else" means a technician, assistant, reprographics or the students. Uses the classes and freezer plan on the All weeks sheet: '
                'until a Year 7 freezer plan is chosen, Weeks 5 to 8 count the prep for both plans.')
    st(pw['A2'], wrap=False)
    S0 = 4
    scols = ['Week', 'Year 7: before the lesson (min)', 'Year 7: after (marking, admin)', 'Year 7: could be done by someone else',
             'Year 8: before the lesson (min)', 'Year 8: after (marking, admin)', 'Year 8: could be done by someone else',
             'Both years, before and after (hours)', 'Optional extras (min, not counted)']
    swidths = [7, 13, 13, 14, 13, 13, 14, 13, 14]
    T0 = S0 + 14  # task table header row
    tcols = ['Week', 'Year', 'When', 'Task', 'Minutes (one class)', 'Minutes from', 'Who could do it', 'Once for all classes?',
             'Reused later', 'Optional?', 'Plan', 'Minutes for your classes']
    twidths = [7, 12, 14, 52, 10, 14, 18, 10, 18, 9, 8, 12]
    for i in range(1, 13):
        pw.column_dimensions[get_column_letter(i)].width = max(swidths[i - 1] if i <= len(swidths) else 0, twidths[i - 1])
    for i, h in enumerate(scols, 1):
        c = pw.cell(row=S0, column=i, value=h); st(c, bold=True, fill=HEAD); c.border = BOX
    r = T0
    for i, h in enumerate(tcols, 1):
        c = pw.cell(row=T0, column=i, value=h); st(c, bold=True, fill=HEAD); c.border = BOX
    for L in lessons:
        for t in sorted(L['prep'], key=lambda t: WHEN_ORDER.index(t['when']) if t['when'] in WHEN_ORDER else 9):
            r += 1
            vals = [L['week'], year_of(L), t['when'], t['task'], t.get('minutes'),
                    {'estimate': 'estimate'}.get(t['minutes_source'], t['minutes_source']), t['who'],
                    'Yes' if t['once_for_all_classes'] else 'No', t.get('reused_later') or None, 'Optional' if t['optional'] else 'No',
                    t['plan'] if t['plan'] != 'Both' else '', None]
            for i, v in enumerate(vals, 1):
                c = pw.cell(row=r, column=i, value=v)
                st(c, align='center' if i in (1, 5, 8, 10, 11, 12) else 'left'); c.border = BOX
                if L['week'] % 2: c.fill = ODD
            if t['minutes_source'] == 'estimate': pw.cell(row=r, column=6).font = Font(name=FONT, size=10, italic=True, color='4D658F')
            pw.cell(row=r, column=12, value=(
                f"=IF(E{r}=\"\",\"\",IF(OR(AND(K{r}=\"Plan A\",'All weeks'!$C$5=\"B\"),AND(K{r}=\"Plan B\",'All weeks'!$C$5=\"A\")),0,"
                f"IF(H{r}=\"Yes\",E{r},E{r}*IF(B{r}=\"Year 7\",'All weeks'!$C$3,'All weeks'!$C$4))))"))
    tl = r
    rng = lambda col: f'${col}${T0 + 1}:${col}${tl}'
    for w in range(1, 11):
        rr = S0 + w
        pw.cell(row=rr, column=1, value=w)
        for yi, (yr, base) in enumerate([('Year 7', 2), ('Year 8', 5)]):
            crit = f'{rng("A")},$A{rr},{rng("B")},"{yr}",{rng("J")},"No"'
            pw.cell(row=rr, column=base, value=f'=SUMIFS({rng("L")},{crit},{rng("C")},"<>After the lesson")')
            pw.cell(row=rr, column=base + 1, value=f'=SUMIFS({rng("L")},{crit},{rng("C")},"After the lesson")')
            pw.cell(row=rr, column=base + 2, value=f'=SUMIFS({rng("L")},{crit},{rng("G")},"<>Teacher")')
        pw.cell(row=rr, column=8, value=f'=ROUND((B{rr}+C{rr}+E{rr}+F{rr})/60,1)')
        pw.cell(row=rr, column=9, value=f'=SUMIFS({rng("L")},{rng("A")},$A{rr},{rng("J")},"Optional")')
        for i in range(1, 10):
            c = pw.cell(row=rr, column=i); st(c, align='center'); c.border = BOX
    rr = S0 + 11
    pw.cell(row=rr, column=1, value='Total')
    for i, col in enumerate('BCDEFGHI', 2):
        pw.cell(row=rr, column=i, value=f'=SUM({col}{S0 + 1}:{col}{S0 + 10})')
    for i in range(1, 10):
        c = pw.cell(row=rr, column=i); st(c, bold=True, align='center', fill=HEAD); c.border = BOX
    for col, colour in (('B', '1F78B4'), ('C', '9DB7D0'), ('E', 'F5B335'), ('F', 'F7D58F'), ('H', '00235F')):
        pw.conditional_formatting.add(f'{col}{S0 + 1}:{col}{S0 + 10}', DataBarRule(start_type='num', start_value=0, end_type='max', color=colour))
    pw.cell(row=T0 - 1, column=1, value='All prep tasks')
    st(pw.cell(row=T0 - 1, column=1), bold=True, wrap=False, size=12)
    pw.freeze_panes = pw.cell(row=T0 + 1, column=1)
    pw.auto_filter.ref = f'A{T0}:L{tl}'
    pw.page_setup.orientation = 'landscape'
    pw.page_setup.paperSize = pw.PAPERSIZE_A4
    pw.page_setup.fitToWidth = 1
    pw.page_setup.fitToHeight = 0
    pw.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------- Cut the prep ----------
    cut_the_prep(cw, lessons, T0, tl, H, last)
    readme(rd, lessons)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print('saved', os.path.abspath(OUT), 'resource rows', last - H, 'prep rows', tl - T0)


def cut_the_prep(ws, lessons, T0, tl, H, last):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 70
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 26
    ws['A1'] = 'Cutting the prep'
    st(ws['A1'], bold=True, size=14, wrap=False)
    ws['A2'] = 'Where the prep time goes for one class, and the options for reducing it. Minutes link to the Prep by week sheet; the lessons themselves are unchanged.'
    st(ws['A2'], wrap=False)
    P = "'Prep by week'!"
    rng = lambda col: f"{P}${col}${T0 + 1}:${col}${tl}"
    A = "'All weeks'!"
    arng = lambda col: f"{A}${col}${H + 1}:${col}${last}"
    r = 4
    blocks = [
        ('1. Do all the printing in one request before Week 1',
         'Send the whole term\'s printing to reprographics at once (sheet totals below, from the All weeks sheet). Printing and sorting then stops being a weekly job.',
         [(f'Printing tasks in the prep list (min, one class)', f'=SUMIFS({rng("E")},{rng("G")},"Reprographics",{rng("J")},"No")')] +
         [(f'{p} (sheets, all weeks)', f'=SUMIFS({arng("L")},{arng("J")},"{p}")') for p in PAPERS]),
        ('2. Hand tasks to a technician, an assistant or the students',
         'Tasks marked as something someone else could do: buying, cutting, counting out kit, setting out trays, clean-up roles.',
         [('Technician or assistant (min, one class)', f'=SUMIFS({rng("E")},{rng("G")},"Technician or assistant",{rng("J")},"No")'),
          ('Students in the lesson (min, one class)', f'=SUMIFS({rng("E")},{rng("G")},"Students (roles)",{rng("J")},"No")'),
          ('Teacher only (min, one class)', f'=SUMIFS({rng("E")},{rng("G")},"Teacher",{rng("J")},"No")')]),
        ('3. Make the reusable sets once, in one session before Week 1',
         'These are made once and reused in later weeks and by every class. Doing them together saves setting up and packing away each time.',
         None),
        ('4. Separate the marking and admin from the prep',
         'Much of the time in Weeks 1, 5 and 10 is marking the skills checks and checkpoints, entering the tracker and putting kit away after the lesson, not getting ready for it.',
         [('After the lesson (min, one class)', f'=SUMIFS({rng("E")},{rng("C")},"After the lesson",{rng("J")},"No")'),
          ('Before the lesson (min, one class)', f'=SUMIFS({rng("E")},{rng("C")},"<>After the lesson",{rng("J")},"No")')]),
        ('5. Drop or simplify the optional extras',
         'Optional tasks are not counted in the totals; skipping them costs nothing in the lesson plan.',
         [('Optional extras (min, one class)', f'=SUMIFS({rng("E")},{rng("J")},"Optional")')]),
    ]
    for title, text, metrics in blocks:
        ws.cell(row=r, column=1, value=title); st(ws.cell(row=r, column=1), bold=True, size=12, wrap=False)
        r += 1
        ws.cell(row=r, column=1, value=text); st(ws.cell(row=r, column=1), wrap=False)
        r += 1
        if metrics:
            for lab, f in metrics:
                ws.cell(row=r, column=3, value=lab); st(ws.cell(row=r, column=3), wrap=False)
                c = ws.cell(row=r, column=4, value=f); st(c, bold=True, align='center'); c.border = BOX
                r += 1
        else:
            for i, h in enumerate(['Week', 'Year', 'Task', 'Minutes', 'Reused'], 1):
                c = ws.cell(row=r, column=i, value=h); st(c, bold=True, fill=HEAD); c.border = BOX
            for L in lessons:
                for t in L['prep']:
                    if t['once_for_all_classes'] and t.get('reused_later') and not t['optional']:
                        r += 1
                        for i, v in enumerate([L['week'], year_of(L), t['task'], t.get('minutes'), t.get('reused_later')], 1):
                            c = ws.cell(row=r, column=i, value=v); st(c, align='center' if i in (1, 4) else 'left'); c.border = BOX
            r += 1
        r += 1
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    # heaviest lessons (computed here for one class, both plans)
    tot = defaultdict(int)
    for L in lessons:
        tot[L['id']] = sum(t['minutes'] or 0 for t in L['prep'] if not t['optional'])
    ws.cell(row=r, column=1, value='6. The heaviest lessons (one class, before and after the lesson, both freezer plans counted)'); st(ws.cell(row=r, column=1), bold=True, size=12, wrap=False)
    r += 1
    for i, h in enumerate(['Week', 'Year', 'Lesson', 'Minutes'], 1):
        c = ws.cell(row=r, column=i, value=h); st(c, bold=True, fill=HEAD); c.border = BOX
    for L in sorted(lessons, key=lambda L: -tot[L['id']])[:6]:
        r += 1
        for i, v in enumerate([L['week'], year_of(L), L['title_en'], tot[L['id']]], 1):
            c = ws.cell(row=r, column=i, value=v); st(c, align='center' if i in (1, 4) else 'left'); c.border = BOX


def readme(ws, lessons):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 130
    flags = [(L, x) for L in lessons for x in L['resources'] if x['in_requisition'] == 'No: add']
    lines = [
        ('Resources and prep, Years 7 and 8, Weeks 1 to 10', True, 14),
        ('All weeks: every resource for every Year 7 and Year 8 lesson on one sheet, grouped by kind (printing, card and stationery, equipment, food and consumables, things to make). Filter by week, year or group. Pink cells are the ones to check.', False, 10),
        ('Prep by week: every prep task with when it happens, how long it takes and who could do it, and a week-by-week total so you can see where the load falls.', False, 10),
        ('Cut the prep: the options for reducing the prep, with the minutes each one saves.', False, 10),
        ('Quantities are for one class of up to 24 students in 6 groups. Type the number of classes and the Year 7 freezer plan in the yellow cells at the top of the All weeks sheet. The "Got it" column is for your ticks.', False, 10),
        ('Checked against every lesson file (lesson plan, kit card, printables), the designed printables, the Teacher Guide prep calendar and the Year 8 Case Bible, and compared with the full requisition workbook. Items already made in an earlier week are listed with "Already have" and are not counted again in the printing totals.', False, 10),
        ('', False, 10),
        (f'Items missing from the full requisition: {len(flags)} rows marked "No: add" on the All weeks sheet.', True, 11),
    ]
    for i, (t, b, s) in enumerate(lines, 1):
        c = ws.cell(row=i + 1, column=2, value=t); st(c, bold=b, size=s)


if __name__ == '__main__':
    build()
